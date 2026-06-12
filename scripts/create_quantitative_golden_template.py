from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config.schema import ROUTE_A_SCHEMA
from config.settings import REPORTS_DIR


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "docs" / "eval" / "quantitative_golden_set_template.xlsx"


def style_header(sheet, row=1):
    for cell in sheet[row]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = f"A{row + 1}"
    sheet.auto_filter.ref = sheet.dimensions


def main():
    workbook = Workbook()
    guide = workbook.active
    guide.title = "填写说明"
    dictionary = workbook.create_sheet("定量指标字典")
    reports = workbook.create_sheet("报告清单")
    entry = workbook.create_sheet("标准答案填写")

    guide.merge_cells("A1:H1")
    guide["A1"] = "A股 ESG 定量指标 Golden Set 填写模板"
    guide["A1"].fill = PatternFill("solid", fgColor="17365D")
    guide["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    guide["A1"].alignment = Alignment(horizontal="center")
    instructions = [
        ("项目", "说明"),
        ("使用方式", "每份报告复制“标准答案填写”中的42行，然后填写黄色列。"),
        ("report_id", "从“报告清单”复制完整报告目录名称。"),
        ("报告原始指标名", "报告中的原始名称与标准指标名不同没有关系。"),
        ("disclosure_status", "disclosed=确认披露；not_found_in_performance_table=绩效表未找到；not_disclosed=确认整份报告未披露；not_checked=未检查。"),
        ("原始值、单位、年份", "填写报告原始内容，不需要提前换算。"),
        ("PDF页码", "填写证据所在 PDF 页码。"),
        ("统计口径/范围", "例如集团口径、境内口径、外购电力等。"),
        ("重要提醒", "绩效表没找到时先填 not_found_in_performance_table，不要直接填 not_disclosed。"),
    ]
    for row_index, values in enumerate(instructions, start=3):
        guide.cell(row_index, 1, values[0])
        guide.cell(row_index, 2, values[1])
        guide.cell(row_index, 1).fill = PatternFill("solid", fgColor="D9EAF7")
        guide.cell(row_index, 1).font = Font(bold=True)
        guide.cell(row_index, 2).alignment = Alignment(wrap_text=True)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 95

    dictionary.append(["field_key", "标准指标名", "分类", "值类型", "单位类型", "常见单位", "别名"])
    for item in ROUTE_A_SCHEMA:
        dictionary.append(
            [
                item["field_key"],
                item["name_cn"],
                item["category"],
                item.get("value_type", ""),
                item.get("unit_type", ""),
                "；".join(item.get("unit_examples", [])),
                "；".join(item.get("aliases", [])),
            ]
        )
    style_header(dictionary)
    for column, width in {"A": 32, "B": 24, "C": 8, "D": 14, "E": 16, "F": 35, "G": 70}.items():
        dictionary.column_dimensions[column].width = width

    reports.append(["report_id", "公司代码", "公司简称", "填写状态"])
    report_names = sorted(
        path.name
        for path in REPORTS_DIR.iterdir()
        if path.is_dir() and path.name != "某个报告目录名"
    )
    for name in report_names:
        parts = name.split("_")
        reports.append([name, parts[0], parts[1] if len(parts) > 1 else "", "未填写"])
    style_header(reports)
    reports.column_dimensions["A"].width = 90
    reports.column_dimensions["B"].width = 14
    reports.column_dimensions["C"].width = 18
    reports.column_dimensions["D"].width = 14
    status_validation = DataValidation(type="list", formula1='"未填写,填写中,已完成"')
    reports.add_data_validation(status_validation)
    status_validation.add(f"D2:D{len(report_names) + 1}")

    headers = [
        "report_id",
        "field_key",
        "标准指标名",
        "分类",
        "报告原始指标名",
        "disclosure_status",
        "原始数值",
        "原始单位",
        "年份",
        "PDF页码",
        "表格标题",
        "统计口径/范围",
        "备注",
    ]
    entry.append(headers)
    for item in ROUTE_A_SCHEMA:
        entry.append(
            [
                "",
                item["field_key"],
                item["name_cn"],
                item["category"],
                "",
                "not_checked",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    style_header(entry)
    for row in entry.iter_rows(min_row=2, max_row=len(ROUTE_A_SCHEMA) + 1, min_col=1, max_col=4):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="E7E6E6")
    for row in entry.iter_rows(min_row=2, max_row=len(ROUTE_A_SCHEMA) + 1, min_col=5, max_col=13):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    disclosure_validation = DataValidation(
        type="list",
        formula1='"disclosed,not_found_in_performance_table,not_disclosed,not_checked"',
    )
    entry.add_data_validation(disclosure_validation)
    disclosure_validation.add(f"F2:F{len(ROUTE_A_SCHEMA) + 1}")
    for column, width in {
        "A": 90,
        "B": 32,
        "C": 24,
        "D": 8,
        "E": 32,
        "F": 35,
        "G": 16,
        "H": 18,
        "I": 12,
        "J": 12,
        "K": 30,
        "L": 35,
        "M": 35,
    }.items():
        entry.column_dimensions[column].width = width

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    print(f"created={OUTPUT_PATH}")
    print(f"quantitative_metrics={len(ROUTE_A_SCHEMA)}")
    print(f"reports={len(report_names)}")


if __name__ == "__main__":
    main()
