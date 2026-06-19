import argparse
import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from config.schema import ROUTE_A_SCHEMA


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT_DIR = PROJECT_ROOT / "docs" / "eval"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "docs" / "eval"

EXPECTED_HEADERS = [
    "report_id",
    "field_key",
    "标准指标名",
    "分类",
    "报告原始指标名",
    "disclosure_status",
    "原始数值",
    "原始单位",
    "年份",
    "PDF页码",
    "表格标题",
    "统计口径/范围",
    "备注",
]

OUTPUT_HEADERS = [
    "source_file",
    "report_id",
    "field_key",
    "standard_metric_name",
    "category",
    "reported_metric_name",
    "disclosure_status",
    "raw_value",
    "raw_unit",
    "year",
    "pdf_page",
    "table_title",
    "scope",
    "notes",
    "status_source",
]

HEADER_MAP = dict(zip(EXPECTED_HEADERS, OUTPUT_HEADERS[1:]))
VALID_STATUSES = {
    "disclosed",
    "not_found_in_performance_table",
    "not_disclosed",
    "not_checked",
    "",
}
TEMPLATE_NAME = "quantitative_golden_set_template.xlsx"
EMPTY_MARKERS = {"", "-", "—", "无", "未披露", "不适用", "na", "n/a", "none", "null"}


def clean_cell(value):
    if value is None:
        return ""
    return str(value).strip()


def is_filled(value):
    return clean_cell(value).lower() not in EMPTY_MARKERS


def infer_disclosure_status(record):
    status = record["disclosure_status"]
    if status and status != "not_checked":
        record["status_source"] = "manual"
        return record

    has_value = is_filled(record["raw_value"])
    has_metric_name = is_filled(record["reported_metric_name"])
    has_evidence_location = is_filled(record["pdf_page"]) or is_filled(record["table_title"])
    if has_value or has_metric_name or has_evidence_location:
        record["disclosure_status"] = "disclosed"
    else:
        record["disclosure_status"] = "not_found_in_performance_table"
    record["status_source"] = "inferred_from_filled_cells"
    return record


def normalize_header(row):
    return [clean_cell(value) for value in row[: len(EXPECTED_HEADERS)]]


def iter_excel_files(input_dir):
    for path in sorted(Path(input_dir).glob("*.xlsx")):
        if path.name.startswith("~$") or path.name == TEMPLATE_NAME:
            continue
        yield path


def read_workbook_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f"{path.name}: empty workbook")

    header = normalize_header(rows[0])
    if header != EXPECTED_HEADERS:
        raise ValueError(
            f"{path.name}: unexpected header {header!r}; expected {EXPECTED_HEADERS!r}"
        )

    records = []
    for row_index, row in enumerate(rows[1:], start=2):
        values = [clean_cell(value) for value in row[: len(EXPECTED_HEADERS)]]
        if not any(values):
            continue
        record = {"source_file": path.name}
        for header_name, value in zip(EXPECTED_HEADERS, values):
            record[HEADER_MAP[header_name]] = value
        record["source_row"] = row_index
        infer_disclosure_status(record)
        records.append(record)
    return records


def build_quality_report(records, input_files):
    schema_fields = {item["field_key"] for item in ROUTE_A_SCHEMA}
    report_ids = sorted({record["report_id"] for record in records if record["report_id"]})
    statuses = Counter(record["disclosure_status"] for record in records)
    duplicate_pairs = [
        f"{report_id}/{field_key}"
        for (report_id, field_key), count in Counter(
            (record["report_id"], record["field_key"]) for record in records
        ).items()
        if count > 1
    ]
    unknown_fields = sorted(
        {record["field_key"] for record in records if record["field_key"] not in schema_fields}
    )
    missing_required = [
        {
            "source_file": record["source_file"],
            "source_row": record["source_row"],
            "missing": [
                key
                for key in ("report_id", "field_key", "disclosure_status")
                if not record.get(key)
            ],
        }
        for record in records
        if not record["report_id"] or not record["field_key"] or not record["disclosure_status"]
    ]
    invalid_statuses = [
        {
            "source_file": record["source_file"],
            "source_row": record["source_row"],
            "status": record["disclosure_status"],
        }
        for record in records
        if record["disclosure_status"] not in VALID_STATUSES
    ]

    expected_records = len(report_ids) * len(schema_fields)
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input_files": [path.name for path in input_files],
        "input_file_count": len(input_files),
        "report_count": len(report_ids),
        "schema_field_count": len(schema_fields),
        "record_count": len(records),
        "expected_record_count": expected_records,
        "complete_grid": len(records) == expected_records,
        "status_counts": dict(sorted(statuses.items())),
        "duplicate_report_fields": duplicate_pairs,
        "unknown_field_keys": unknown_fields,
        "missing_required_cells": missing_required,
        "invalid_statuses": invalid_statuses,
    }


def write_csv(path, records):
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = OUTPUT_HEADERS + ["source_row"]
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_quality_markdown(path, quality):
    lines = [
        "# Quantitative Golden Set Quality Report",
        "",
        f"- Generated at: `{quality['generated_at']}`",
        f"- Input Excel files: {quality['input_file_count']}",
        f"- Reports: {quality['report_count']}",
        f"- Schema fields: {quality['schema_field_count']}",
        f"- Records: {quality['record_count']}",
        f"- Expected records: {quality['expected_record_count']}",
        f"- Complete report-field grid: `{quality['complete_grid']}`",
        "",
        "## Disclosure Status Counts",
        "",
    ]
    for status, count in quality["status_counts"].items():
        lines.append(f"- `{status or '<blank>'}`: {count}")

    checks = [
        ("Duplicate Report Fields", "duplicate_report_fields"),
        ("Unknown Field Keys", "unknown_field_keys"),
        ("Missing Required Cells", "missing_required_cells"),
        ("Invalid Statuses", "invalid_statuses"),
    ]
    for title, key in checks:
        values = quality[key]
        lines.extend(["", f"## {title}", ""])
        if not values:
            lines.append("- None")
        else:
            for value in values[:100]:
                lines.append(f"- `{value}`")
            if len(values) > 100:
                lines.append(f"- ... {len(values) - 100} more")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_golden_set(input_dir=DEFAULT_INPUT_DIR, output_dir=DEFAULT_OUTPUT_DIR):
    input_files = list(iter_excel_files(input_dir))
    records = []
    for path in input_files:
        records.extend(read_workbook_rows(path))

    quality = build_quality_report(records, input_files)
    output_dir = Path(output_dir)
    write_csv(output_dir / "quantitative_golden_set.csv", records)
    write_json(output_dir / "quantitative_golden_set.json", records)
    write_json(output_dir / "quantitative_golden_set_quality.json", quality)
    write_quality_markdown(output_dir / "quantitative_golden_set_quality.md", quality)
    return quality


def parse_args():
    parser = argparse.ArgumentParser(
        description="Merge per-report quantitative Golden Set Excel files."
    )
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def main():
    args = parse_args()
    quality = build_golden_set(args.input_dir, args.output_dir)
    print(f"input_files={quality['input_file_count']}")
    print(f"reports={quality['report_count']}")
    print(f"records={quality['record_count']}")
    print(f"complete_grid={quality['complete_grid']}")
    if quality["invalid_statuses"] or quality["unknown_field_keys"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
